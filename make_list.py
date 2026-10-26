import os
import re
import json
import shutil
from datetime import datetime
from common import load_config
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image
from concurrent.futures import ThreadPoolExecutor, as_completed

Image.MAX_IMAGE_PIXELS = None

VALID_EXTENSIONS = (".emf", ".png", ".jpg", ".jpeg", ".bmp")
PNG_STAMP_NAME = ".convert_settings.json"
OTHER_KEY = 999


# ファイル名の空白・連続アンダーバーを整理する（"sp_1_g_My_ST_ EX" → "sp_1_g_My_ST_EX"）
def sanitize_filename(filename):
    name = os.path.splitext(filename)[0]
    name = name.replace(' ', '')
    name = re.sub(r'_+', '_', name).strip('_')
    return name


# -----------------------------------------------------------------------------
# 画像変換 (EMF 等 → PNG)
# -----------------------------------------------------------------------------
def _convert_single_file(filename, image_folder, png_folder, max_width, colors, force):
    src_path = os.path.join(image_folder, filename)
    png_path = os.path.join(png_folder, sanitize_filename(filename) + ".png")

    if os.path.exists(png_path) and not force:
        return "skipped", filename
    try:
        with Image.open(src_path) as img:
            img.load()
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                background = Image.new("RGB", img.size, "white")
                background.paste(img, mask=img.getchannel("A"))
                img = background
            else:
                img = img.convert("RGB")
            orig_w, orig_h = img.size
            if orig_w > max_width:
                new_h = int(orig_h * (max_width / orig_w))
                img = img.resize((max_width, new_h), Image.Resampling.LANCZOS)
            if colors > 0:
                img = img.quantize(colors, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
            img.save(png_path, "PNG", optimize=True)
        return "success", filename
    except Exception as e:
        return "error", f"{filename} (- {e})"


def convert_to_png(image_folder, png_folder, max_width, colors):
    src_files = [f for f in os.listdir(image_folder) if f.lower().endswith(VALID_EXTENSIONS)]
    total_files = len(src_files)
    if total_files == 0:
        return

    # 変換条件（幅・色数）を変えたときは、変換済み PNG も作り直す
    stamp_path = os.path.join(png_folder, PNG_STAMP_NAME)
    settings = {"max_width": max_width, "colors": colors}
    try:
        with open(stamp_path, "r", encoding="utf-8") as f:
            force = json.load(f) != settings
    except (OSError, ValueError):
        force = os.path.exists(stamp_path)

    print("=" * 70)
    print(f" 画像の PNG への自動変換・軽量化を開始します（全 {total_files} 件）...")
    print(f" 保存先フォルダ: {png_folder}")
    print(f" 最大幅: {max_width}px / 色数: {colors if colors > 0 else 'フルカラー'}"
          + (" （設定変更のため再変換）" if force else ""))
    print("=" * 70)

    converted, skipped, failed = 0, 0, 0
    # EMF は巨大なビットマップに展開されメモリを使うため、並列数は控えめにする
    max_workers = min(os.cpu_count() or 4, 4)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(_convert_single_file, filename, image_folder, png_folder,
                            max_width, colors, force): filename
            for filename in src_files
        }
        for i, future in enumerate(as_completed(futures), 1):
            status, info = future.result()
            if status == "success": converted += 1
            elif status == "skipped": skipped += 1
            else:
                failed += 1
                print(f"\n[エラー] 変換失敗: {info}")

            percent = int((i / total_files) * 100)
            print(
                f"\r 変換進行中: [{percent:3d}%] ({i}/{total_files}枚完了)  "
                f"[新規変換:{converted} / スキップ:{skipped} / 失敗:{failed}]      ",
                end="", flush=True
            )
    print("\n" + "=" * 70)

    if failed == 0:
        with open(stamp_path, "w", encoding="utf-8") as f:
            json.dump(settings, f)


# -----------------------------------------------------------------------------
# ファイル名の解析（config.yaml の filename_rules に従う）
# -----------------------------------------------------------------------------
def _index_or_other(value, order, ignore_case=False):
    """order 内での位置。value が無ければ OTHER_KEY、order に無ければ末尾扱い。"""
    if value is None:
        return OTHER_KEY
    for i, item in enumerate(order):
        a, b = str(item), value
        if (a.lower() == b.lower()) if ignore_case else (a == b):
            return i
    return len(order)


def _lookup(mapping, key, ignore_case=True):
    """変換表から値を引く。見つからなければキー自身を返す。"""
    if key is None:
        return ""
    if key in mapping:
        return mapping[key]
    if ignore_case:
        for k, v in mapping.items():
            if str(k).lower() == key.lower():
                return v
    return key


def parse_filename(filename, config):
    """戻り値: (ソートキー, 画像タイトル, 章タイトル, 単独ページ"0"/"1")"""
    base_name = sanitize_filename(filename)

    floors_order = config.get("sort_order_floors", [])
    members_order = config.get("sort_order_members", [])
    stresses_order = config.get("sort_order_stresses", [])
    cases_order = config.get("sort_order_load_cases", [])

    for rule in config.get("filename_rules", []):
        m = re.match(rule["pattern"], base_name)
        if not m:
            continue
        g = m.groupdict()
        floor, member, stress, case = g.get("floor"), g.get("member"), g.get("stress"), g.get("case")

        # 応力記号が sort_order_stresses に無いものは、このルールに該当しないとみなす
        if stress is not None and stress not in stresses_order:
            continue

        floor_key = _index_or_other(floor, floors_order)
        member_key = _index_or_other(member, members_order, True)
        stress_key = _index_or_other(stress, stresses_order)
        case_key = _index_or_other(case, cases_order, True)
        type_key = int(rule.get("type_order", 1))

        if "order_from_case" in rule:
            chap_order = int(rule["order_from_case"]) + case_key
        else:
            chap_order = int(rule["order"])

        if rule.get("sort_by") == "type":
            sort_key = (chap_order, 0, type_key, floor_key, member_key, stress_key)
        else:
            sort_key = (chap_order, member_key, floor_key, stress_key, case_key, type_key)

        case_ja = _lookup(config.get("title_map_load_cases", {}), case)
        values = {
            "floor": floor or "",
            "floor_ja": "" if floor is None else (floor if "階" in floor else f"{floor}階"),
            "member": member or "",
            "member_ja": _lookup(config.get("title_map_members", {}), member),
            "stress": stress or "",
            "stress_ja": _lookup(config.get("title_map_stresses", {}), stress, False),
            "case": case or "",
            "case_ja": case_ja,
        }
        try:
            title = rule["title"].format(**values)
        except KeyError:
            title = base_name

        if rule.get("chapter_by") == "case":
            chapter = _lookup(config.get("title_map_chapters", {}), case)
            if chapter == case:
                chapter = f"応力図 ({case_ja})"
        else:
            chapter = config.get("title_map_sections", {}).get(rule.get("section"), rule.get("name", "その他"))

        single = bool(rule.get("single_page", False))
        for group_name, exc_values in (rule.get("single_page_exceptions") or {}).items():
            target = g.get(group_name)
            if target is not None and any(target.lower() == str(v).lower() for v in exc_values):
                single = not single
                break

        return (sort_key, title, chapter, "1" if single else "0")

    return ((OTHER_KEY,) * 6, base_name, "その他", "0")


# -----------------------------------------------------------------------------
# リスト (Excel) 作成
# -----------------------------------------------------------------------------
def backup_existing_list(list_path, list_filename, base_dir):
    if not os.path.exists(list_path):
        return
    backup_dir = os.path.join(base_dir, "list_backup")
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"{os.path.splitext(list_filename)[0]}_{stamp}.xlsx")
    try:
        shutil.copy2(list_path, backup_path)
        print(f"既存のリストをバックアップしました -> {os.path.relpath(backup_path, base_dir)}")
    except OSError as e:
        print(f"Warning: バックアップに失敗しました: {e}")


def write_list_xlsx(rows, list_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "list"
    ws.append(["画像ファイル名", "画像タイトル", "章タイトル", "単独ページ"])
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCD", (42, 50, 36, 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(list_path)


def generate_image_list():
    config, base_dir = load_config()
    image_folder = os.path.join(base_dir, config.get("image_folder", "images"))

    png_folder = os.path.join(base_dir, "PNG")
    os.makedirs(png_folder, exist_ok=True)

    list_filename = config.get("list_path", "list.xlsx")
    list_path = os.path.join(base_dir, list_filename)

    if not os.path.exists(image_folder):
        print(f"Error: Target image directory does not exist: {image_folder}")
        return

    convert_to_png(image_folder, png_folder,
                   int(config.get("png_max_width_px", 3000)),
                   int(config.get("png_colors", 256)))

    parsed_files = []
    for f in os.listdir(image_folder):
        if f.lower().endswith(VALID_EXTENSIONS):
            sort_key, title, chapter, single = parse_filename(f, config)
            out_img_name = "PNG/" + sanitize_filename(f) + ".png"
            parsed_files.append((sort_key, out_img_name, title, chapter, single))

    if not parsed_files:
        print("Warning: No valid image files found.")
        return

    # 同順位はファイル名順にして、実行のたびに並びが変わらないようにする
    parsed_files.sort(key=lambda x: (x[0], x[1]))

    seen_chapters = set()
    rows = []
    for _, img, title, chapter, single in parsed_files:
        # 章タイトルは各章の最初の行にだけ記載し、以降は "-" とする
        if chapter not in seen_chapters:
            seen_chapters.add(chapter)
            chap_cell = chapter
        else:
            chap_cell = "-"
        rows.append((img, title, chap_cell, single))

    backup_existing_list(list_path, list_filename, base_dir)
    try:
        write_list_xlsx(rows, list_path)
    except PermissionError:
        print(f"Error: {list_filename} が Excel で開かれています。閉じてから再実行してください。")
        return

    others = sum(1 for r in parsed_files if r[3] == "その他")
    print("-" * 70)
    print(f"Success: リストファイルを作成しました -> {list_filename}")
    if others:
        print(f"Warning: ルールに一致しない画像が {others} 件あります（章タイトル「その他」）。"
              f"config.yaml の filename_rules を確認してください。")
    print("-" * 70)


if __name__ == "__main__":
    generate_image_list()
    input("\nEnterキーを押して終了してください...")
