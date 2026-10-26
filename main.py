import os
import sys
import logging
import io
from common import load_config, get_base_dir
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    handlers=[
        logging.FileHandler(os.path.join(get_base_dir(), "process.log"), encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)

def set_font_style(run, font_name, size_pt):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)

def add_field(paragraph, instr, placeholder, font_name, size_pt, dirty=False):
    """Word のフィールド（SEQ など）を段落に追加する。placeholder は更新前の仮表示。"""
    def _run():
        run = paragraph.add_run()
        set_font_style(run, font_name, size_pt)
        return run

    begin = _run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    if dirty:
        fld_begin.set(qn('w:dirty'), 'true')
    begin._r.append(fld_begin)

    instr_run = _run()
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = f' {instr} '
    instr_run._r.append(instr_text)

    sep = _run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    sep._r.append(fld_sep)

    text = _run()
    text.text = placeholder

    end = _run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    end._r.append(fld_end)


def style_or_none(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def setup_style_font(style, font_name, size_pt, bold=None):
    """見出し・図表番号スタイルの書式を設定する（テンプレート使用時は呼ばない）。"""
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rFonts.set(qn(attr), font_name)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme'):
        if rFonts.get(qn(attr)) is not None:
            del rFonts.attrib[qn(attr)]


LIST_COLUMNS = ['画像ファイル名', '画像タイトル', '章タイトル', '単独ページ']


def read_list_rows(list_path):
    """list.xlsx を読み、[(Excel の行番号, {列名: 文字列})] を返す。必須列が無ければ ValueError。"""
    wb = load_workbook(list_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        data = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(data, ())]
        missing = [c for c in LIST_COLUMNS if c not in header]
        if missing:
            raise ValueError(f"必須列がありません: {', '.join(missing)}")
        idx = {c: header.index(c) for c in LIST_COLUMNS}
        rows = []
        for row_num, values in enumerate(data, start=2):
            rows.append((row_num, {
                c: ("" if idx[c] >= len(values) or values[idx[c]] is None else str(values[idx[c]]).strip())
                for c in LIST_COLUMNS
            }))
        return rows
    finally:
        wb.close()


def create_professional_report():
    config, base_dir = load_config()

    list_filename = config.get("list_path", "list.xlsx")
    list_path = os.path.join(base_dir, list_filename)

    output_word_str = str(config.get("output_word_path", "output.docx")).strip()
    output_word_path = os.path.isabs(output_word_str) and output_word_str or os.path.join(base_dir, output_word_str)

    # 白黒変換フラグの取得（デフォルトは False）
    convert_to_grayscale = config.get("convert_to_grayscale", False)

    if not os.path.exists(list_path):
        logging.error(f"Dataset file not found: {list_path}. Please run make_list.py first.")
        return

    try:
        rows = read_list_rows(list_path)
        logging.info(f"Loaded source file: [{os.path.basename(list_path)}]")
    except Exception as e:
        logging.error(f"Failed to read dataset file: {e}")
        return

    tmpl = config.get("template_word_path", "")
    doc = None
    tmpl_loaded = False
    if tmpl and str(tmpl).strip() and tmpl != "None":
        tmpl_path = os.path.isabs(str(tmpl).strip()) and str(tmpl).strip() or os.path.join(base_dir, str(tmpl).strip())
        if os.path.exists(tmpl_path):
            try:
                doc = Document(tmpl_path)
                tmpl_loaded = True
                logging.info(f"Using template Word document: {tmpl_path}")
            except Exception as t_err:
                logging.warning(f"Failed to load template, fallback to blank doc: {t_err}")
                doc = Document()
        else:
            doc = Document()
    else:
        doc = Document()

    font_family = str(config.get("font_family", "MS Pゴシック"))
    try:
        chap_title_pt = float(config.get("chap_title_pt", 12))
        img_title_pt = float(config.get("img_title_pt", 9))
        max_height_mm = float(config.get("max_img_height_mm", 90))
        margin_mm = float(config.get("page_margin_mm", 20))
    except ValueError:
        chap_title_pt, img_title_pt, max_height_mm, margin_mm = 12.0, 9.0, 90.0, 20.0

    for section in doc.sections:
        section.top_margin = Mm(margin_mm)
        section.bottom_margin = Mm(margin_mm)
        section.left_margin = Mm(margin_mm)
        section.right_margin = Mm(margin_mm)

    chapter_in_body = str(config.get("chapter_title_position", "header")).lower() == "body"
    figure_numbering = bool(config.get("figure_numbering", False))
    figure_label = str(config.get("figure_label", "図"))
    page_break_keywords = list(config.get("page_break_keywords", []))
    heading_style = style_or_none(doc, "Heading 1") if chapter_in_body else None
    caption_style = style_or_none(doc, "Caption") if figure_numbering else None
    if not tmpl_loaded:
        if heading_style is not None:
            setup_style_font(heading_style, font_family, chap_title_pt, bold=True)
        if caption_style is not None:
            setup_style_font(caption_style, font_family, img_title_pt, bold=False)
    figure_count = 0

    section = doc.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin

    images_on_current_page = 0
    last_processed_chap = ""
    last_processed_type = ""
    is_first_section = True

    for current_row_num, row in rows:
        try:
            f_name = row['画像ファイル名']
            img_title = row['画像タイトル']
            chap_title = row['章タイトル']
            single_val = row['単独ページ']

            if chap_title in ["-", "なし", "nan", "None", "none", ""]:
                chap_title = ""
            if img_title in ["-", "なし", "nan", "None", "none", ""]:
                img_title = ""

            if img_title: img_title = img_title.replace("_", " ")
            if chap_title: chap_title = chap_title.replace("_", " ")

            is_single = single_val in ["1", "1.0", "True", "true"]

            if not f_name or f_name in ["-", "なし", "nan"]:
                continue

            # 画像ファイルのフルパス解決
            img_full_path = os.path.join(base_dir, f_name)
            if not os.path.exists(img_full_path):
                logging.warning(f"Row {current_row_num}: Image missing [{f_name}]")
                continue

            # --- 章見出し（セクションヘッダー）の処理 ---
            if chap_title and chap_title != last_processed_chap:
                if not is_first_section:
                    new_section = doc.add_section(2)
                    new_section.header.is_linked_to_previous = False
                    new_section.first_page_header.is_linked_to_previous = False
                    current_section = new_section
                else:
                    current_section = doc.sections[0]
                    is_first_section = False

                if chapter_in_body:
                    p_head = doc.add_paragraph(style=heading_style)
                    run_head = p_head.add_run(chap_title)
                    run_head.bold = True
                    set_font_style(run_head, font_family, chap_title_pt)
                else:
                    current_section.different_first_page_header_footer = True
                    p_head = current_section.first_page_header.paragraphs[0]
                    p_head.text = ""
                    run_head = p_head.add_run(chap_title)
                    run_head.bold = True
                    set_font_style(run_head, font_family, chap_title_pt)

                    current_section.header.paragraphs[0].text = ""
                last_processed_chap = chap_title
                images_on_current_page = 0
                last_processed_type = ""

            # --- 自動改ページ処理（種類・部材の切り替え時） ---
            current_type = ""
            for kw in page_break_keywords:
                if kw in img_title:
                    current_type = kw
                    break

            if not current_type:
                for m_ja in config.get("title_map_members", {}).values():
                    if m_ja in img_title:
                        current_type = m_ja
                        break

            if current_type and last_processed_type and current_type != last_processed_type:
                doc.add_page_break()
                images_on_current_page = 0

            if current_type:
                last_processed_type = current_type

            if images_on_current_page >= 2:
                doc.add_page_break()
                images_on_current_page = 0

            # --- 画像のサイズ計算・白黒変換・挿入 ---
            with Image.open(img_full_path) as img_file:
                orig_width, orig_height = img_file.size
                img_byte_arr = io.BytesIO()
                if convert_to_grayscale:
                    # 白黒化したものを PNG で埋め込む
                    img_file.convert('L').save(img_byte_arr, format='PNG', optimize=True)
                else:
                    # 再エンコードせず、PNG のバイト列をそのまま埋め込む（劣化・肥大化を防ぐ）
                    with open(img_full_path, "rb") as raw:
                        img_byte_arr.write(raw.read())
                img_byte_arr.seek(0)

            aspect_ratio = orig_height / orig_width

            try:
                single_width = float(config.get("single_img_width_mm", 150))
                default_width = float(config.get("default_img_width_mm", 120))
            except ValueError:
                single_width, default_width = 150.0, 120.0

            selected_width_mm = single_width if is_single else default_width
            target_width_mm = selected_width_mm * float(config.get("img_width_scale", 1.15))
            expected_height_mm = target_width_mm * aspect_ratio

            if expected_height_mm > max_height_mm:
                expected_height_mm = max_height_mm
                target_width_mm = expected_height_mm / aspect_ratio

            target_width = Mm(target_width_mm)
            if target_width > max_width:
                target_width = max_width

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(6)

            run_img = p_img.add_run()
            # パスではなくメモリ上の画像データを挿入
            run_img.add_picture(img_byte_arr, width=target_width)

            if img_title:
                p_title = doc.add_paragraph(style=caption_style)
                p_title.paragraph_format.space_before = Pt(4)
                p_title.paragraph_format.space_after = Pt(18)
                p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if figure_numbering:
                    figure_count += 1
                    run_label = p_title.add_run(f"{figure_label} ")
                    set_font_style(run_label, font_family, img_title_pt)
                    add_field(p_title, f"SEQ {figure_label} \\* ARABIC", str(figure_count),
                              font_family, img_title_pt)
                    run_title = p_title.add_run(f"  {img_title}")
                else:
                    run_title = p_title.add_run(img_title)
                set_font_style(run_title, font_family, img_title_pt)
            else:
                p_img.paragraph_format.space_after = Pt(18)

            if is_single:
                images_on_current_page = 2
            else:
                images_on_current_page += 1

        except Exception as row_err:
            logging.error(f"Row {current_row_num}: Exception encountered: {row_err}")
            continue

    # --- 保存とPDF自動変換処理 ---
    try:
        doc.save(output_word_path)
        logging.info(f"Success: Wordファイルを出力しました -> {output_word_path}")

        try:
            from docx2pdf import convert
            output_pdf_path = os.path.splitext(output_word_path)[0] + ".pdf"
            logging.info("WordファイルをPDFに自動変換しています...")
            convert(output_word_path, output_pdf_path)
            logging.info(f"Success: PDFファイルを出力しました -> {output_pdf_path}")
        except ImportError:
            logging.warning("docx2pdf ライブラリがないためPDF変換をスキップしました (pip install docx2pdf で有効になります)")
        except Exception as pdf_err:
            logging.error(f"PDF変換中にエラーが発生しました: {pdf_err}")

    except Exception as save_exc:
        logging.error(f"I/O Error: Failed to write output file: {save_exc}")

if __name__ == "__main__":
    create_professional_report()
    input("\nProcess finished. Press Enter to exit...")